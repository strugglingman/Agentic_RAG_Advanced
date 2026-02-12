"""
Shared file content extraction utilities.

Used by both AgentService (ReAct) and LangGraph nodes (Plan-Execute)
to extract text content from file attachments.

Supported formats:
    - Images: image/png, image/jpeg, image/gif, image/webp (via Vision API)
    - PDF: application/pdf
    - DOCX: application/vnd.openxmlformats-officedocument.wordprocessingml.document
    - Excel: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
    - Text: text/plain, text/markdown, text/csv
"""

import base64
import logging
import os

from src.services.llm_client import chat_completion
from src.config.settings import Config

logger = logging.getLogger(__name__)

MAX_CONTENT_CHARS = 50000  # Limit to prevent overwhelming context


async def describe_image_with_vision(
    file_path: str, mime_type: str, openai_client=None
) -> str:
    """
    Use Vision API to describe an image.

    Args:
        file_path: Absolute path to image file
        mime_type: MIME type (image/png, image/jpeg, etc.)
        openai_client: OpenAI client instance

    Returns:
        Text description of the image from Vision API
    """
    if not openai_client:
        file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0
        return f"[Image file ({mime_type}) - {file_size} bytes - Vision API client not available]"

    try:
        # Read and encode image as base64
        with open(file_path, "rb") as f:
            image_data = base64.b64encode(f.read()).decode("utf-8")

        # Build multimodal message for Vision API
        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Describe this image in detail. Include:\n"
                            "- What the image shows (objects, people, scenes)\n"
                            "- Any text visible in the image\n"
                            "- Charts/graphs: describe the data and trends\n"
                            "- Documents: summarize the content\n"
                            "Be concise but thorough."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{mime_type};base64,{image_data}",
                            "detail": "auto",
                        },
                    },
                ],
            }
        ]

        # Call Vision API
        response = await chat_completion(
            client=openai_client,
            model=Config.OPENAI_VISION_MODEL,
            messages=messages,
            temperature=0.3,
            max_tokens=1000,
        )

        description = response.choices[0].message.content.strip()
        logger.info(f"Vision API described image: {file_path[:50]}...")
        return f"[IMAGE DESCRIPTION]\n{description}"

    except Exception as e:
        logger.error(f"Vision API failed for {file_path}: {e}")
        try:
            file_size = os.path.getsize(file_path)
            return f"[Image file ({mime_type}) - {file_size} bytes - Vision API error: {e}]"
        except Exception:
            return f"[Image file ({mime_type}) - Vision API unavailable]"


async def extract_file_content(
    file_path: str, mime_type: str, openai_client=None
) -> str:
    """
    Extract text content from file based on MIME type.

    Args:
        file_path: Absolute path to file on disk
        mime_type: MIME type of the file
        openai_client: Optional OpenAI client for Vision API (images)

    Returns:
        Extracted text content (truncated if too long)
    """
    try:
        # Image files - use Vision API to describe
        if mime_type.startswith("image/"):
            return await describe_image_with_vision(file_path, mime_type, openai_client)

        # PDF files
        elif mime_type == "application/pdf":
            try:
                import PyPDF2

                with open(file_path, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    text_parts = []
                    for page in reader.pages[:50]:  # Limit to first 50 pages
                        text_parts.append(page.extract_text())
                    content = "\n".join(text_parts)
                    return content[:MAX_CONTENT_CHARS] + (
                        "..." if len(content) > MAX_CONTENT_CHARS else ""
                    )
            except Exception as e:
                logger.error(f"Failed to extract PDF content: {e}")
                return f"[PDF file - text extraction failed: {e}]"

        # DOCX files
        elif (
            mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            try:
                import docx

                doc = docx.Document(file_path)
                text_parts = [paragraph.text for paragraph in doc.paragraphs]
                content = "\n".join(text_parts)
                return content[:MAX_CONTENT_CHARS] + (
                    "..." if len(content) > MAX_CONTENT_CHARS else ""
                )
            except Exception as e:
                logger.error(f"Failed to extract DOCX content: {e}")
                return f"[DOCX file - text extraction failed]"

        # Excel files
        elif mime_type in [
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ]:
            try:
                import openpyxl

                wb = openpyxl.load_workbook(file_path, data_only=True)
                text_parts = []
                for sheet in wb.worksheets[:5]:  # Limit to first 5 sheets
                    text_parts.append(f"Sheet: {sheet.title}")
                    for row in list(sheet.iter_rows(values_only=True))[
                        :100
                    ]:  # Limit to 100 rows
                        row_text = "\t".join(
                            str(cell) if cell is not None else "" for cell in row
                        )
                        if row_text.strip():
                            text_parts.append(row_text)
                content = "\n".join(text_parts)
                return content[:MAX_CONTENT_CHARS] + (
                    "..." if len(content) > MAX_CONTENT_CHARS else ""
                )
            except Exception as e:
                logger.error(f"Failed to extract Excel content: {e}")
                return f"[Excel file - text extraction failed]"

        # Text files (plain text, markdown, CSV, etc.)
        elif mime_type.startswith("text/"):
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
                    return content[:MAX_CONTENT_CHARS] + (
                        "..." if len(content) > MAX_CONTENT_CHARS else ""
                    )
            except UnicodeDecodeError:
                try:
                    with open(file_path, "r", encoding="latin-1") as f:
                        content = f.read()
                        return content[:MAX_CONTENT_CHARS] + (
                            "..." if len(content) > MAX_CONTENT_CHARS else ""
                        )
                except Exception as e:
                    logger.error(f"Failed to read text file: {e}")
                    return f"[Text file - encoding error]"

        # Unsupported file type
        else:
            file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0
            return f"[File type {mime_type} not supported - {file_size} bytes]"

    except Exception as e:
        logger.error(f"Unexpected error extracting file content: {e}")
        return f"[Error reading file: {str(e)}]"
